import re
from decimal import Decimal
from pathlib import Path

import openpyxl

from .errors import CardValidationError
from .models import CardRecord, SalePlanRow

TRACKER_SHEET = "Tracker"
SALE_PLAN_SHEET = "Sale Plan"
TRACKER_HEADER_ROW = 4

# Tracker sheet header text -> CardRecord field name.
REQUIRED_COLUMNS = {
    "ID": "card_id",
    "Card Name": "character",
    "Set / Series": "set_name",
    "Card No.": "card_number",
    "Rarity / Variant": "rarity",
    "Language": "language",
    "Grade": "graded",
    "Notes": "notes",
}
TYPE_COLUMN = "Type"
CONDITION_COLUMN = "Condition"

BUNDLE_TAG_PATTERN = re.compile(r"^\[(.*?)\]")


def _card_id(cell_value) -> str:
    if isinstance(cell_value, (int, float)):
        return str(int(cell_value))
    return str(cell_value).strip()


def read_inventory(xlsx_path: str | Path) -> list[CardRecord]:
    workbook = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=True)
    if TRACKER_SHEET not in workbook.sheetnames:
        raise CardValidationError(row=0, column=TRACKER_SHEET, message=f"workbook has no sheet named '{TRACKER_SHEET}'")
    sheet = workbook[TRACKER_SHEET]
    rows = sheet.iter_rows(min_row=TRACKER_HEADER_ROW, values_only=True)

    try:
        header = next(rows)
    except StopIteration:
        raise CardValidationError(row=TRACKER_HEADER_ROW, column="<header>", message="sheet is empty, no header row found")

    column_index: dict[str, int] = {}
    for expected_header in list(REQUIRED_COLUMNS) + [TYPE_COLUMN, CONDITION_COLUMN]:
        try:
            column_index[expected_header] = header.index(expected_header)
        except ValueError:
            if expected_header in (TYPE_COLUMN, CONDITION_COLUMN):
                column_index[expected_header] = -1  # optional, tolerate absence
                continue
            raise CardValidationError(row=TRACKER_HEADER_ROW, column=expected_header, message="required column missing from header row")

    records: list[CardRecord] = []
    seen_at_row: dict[str, int] = {}
    for offset, row in enumerate(rows, start=1):
        row_number = TRACKER_HEADER_ROW + offset
        if row is None or all(cell is None for cell in row):
            continue

        values: dict[str, str] = {}
        for expected_header, field_name in REQUIRED_COLUMNS.items():
            idx = column_index[expected_header]
            cell_value = row[idx] if idx < len(row) else None
            if cell_value is None or (isinstance(cell_value, str) and not cell_value.strip()):
                raise CardValidationError(row=row_number, column=expected_header, message="required value is missing")
            values[field_name] = _card_id(cell_value) if field_name == "card_id" else str(cell_value).strip()

        notes = values.pop("notes")
        tag_match = BUNDLE_TAG_PATTERN.match(notes)
        if not tag_match:
            raise CardValidationError(row=row_number, column="Notes", message=f"expected a leading bracket tag like '[SINGLE]' or '[BUNDLE X - ...]', got: {notes!r}")
        bundle_tag = tag_match.group(1)

        type_idx = column_index[TYPE_COLUMN]
        if type_idx != -1 and type_idx < len(row) and row[type_idx] is not None:
            type_value = str(row[type_idx]).strip()
            expected_type = "Single" if bundle_tag == "SINGLE" else "Lot"
            if type_value != expected_type:
                raise CardValidationError(row=row_number, column=TYPE_COLUMN, message=f"'{type_value}' does not match Notes tag '[{bundle_tag}]' (expected '{expected_type}')")

        condition_idx = column_index[CONDITION_COLUMN]
        condition_value = row[condition_idx] if condition_idx != -1 and condition_idx < len(row) else None
        condition = str(condition_value).strip() if condition_value is not None and str(condition_value).strip() else None

        record = CardRecord(bundle_tag=bundle_tag, condition=condition, **values)
        if record.card_id in seen_at_row:
            raise CardValidationError(row=row_number, column="ID", message=f"duplicate card_id, first seen at row {seen_at_row[record.card_id]}")
        seen_at_row[record.card_id] = row_number
        records.append(record)

    return records


def read_sale_plan(xlsx_path: str | Path) -> list[SalePlanRow]:
    workbook = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=True)
    if SALE_PLAN_SHEET not in workbook.sheetnames:
        raise CardValidationError(row=0, column=SALE_PLAN_SHEET, message=f"workbook has no sheet named '{SALE_PLAN_SHEET}'")
    sheet = workbook[SALE_PLAN_SHEET]

    rows: list[SalePlanRow] = []
    for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        num = row[1] if len(row) > 1 else None
        if not isinstance(num, (int, float)):
            continue  # skip title/header/blank rows
        listing_text = row[2] if len(row) > 2 else None
        card_count = row[3] if len(row) > 3 else None
        list_at = row[4] if len(row) > 4 else None
        floor = row[5] if len(row) > 5 else None

        if not listing_text or not str(listing_text).strip():
            raise CardValidationError(row=row_number, column="Listing", message="required value is missing")
        if list_at is None:
            raise CardValidationError(row=row_number, column="List at", message="required value is missing")
        if floor is None:
            raise CardValidationError(row=row_number, column="Floor", message="required value is missing")

        rows.append(
            SalePlanRow(
                number=int(num),
                listing_text=str(listing_text).strip(),
                card_count=int(card_count) if card_count is not None else 0,
                list_price=Decimal(str(list_at)),
                floor_price=Decimal(str(floor)),
            )
        )
    return rows
