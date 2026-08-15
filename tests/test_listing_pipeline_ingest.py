import openpyxl
import pytest

from automation_control.listing_pipeline.errors import CardValidationError
from automation_control.listing_pipeline.ingest import read_inventory, read_sale_plan

TRACKER_HEADER = ["ID", "Card Name", "Set / Series", "Card No.", "Rarity / Variant", "Language", "Grade", "Condition", "Type", "Notes"]


def _write_tracker(path, rows):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Tracker"
    sheet.append([])
    sheet.append(["Card Listing Tracker"])
    sheet.append(["One row per card"])
    sheet.append(TRACKER_HEADER)
    for row in rows:
        sheet.append(row)
    workbook.save(path)


def _tracker_row(card_id, name, set_name="Base Set", card_no="1/1", rarity="Rare", language="English", grade="Raw", condition=None, card_type="Single", notes="[SINGLE] note"):
    return [card_id, name, set_name, card_no, rarity, language, grade, condition, card_type, notes]


def test_valid_rows_ingest_correctly(tmp_path):
    path = tmp_path / "tracker.xlsx"
    _write_tracker(path, [
        _tracker_row(1, "Pikachu", card_no="58/102"),
        _tracker_row(2, "Charizard", card_no="4/102", card_type="Lot", notes="[BUNDLE A - Test lot] note"),
    ])
    cards = read_inventory(path)
    assert [card.card_id for card in cards] == ["1", "2"]
    assert cards[0].character == "Pikachu"
    assert cards[0].bundle_tag == "SINGLE"
    assert cards[0].condition is None
    assert cards[1].bundle_tag == "BUNDLE A - Test lot"


def test_condition_is_optional_but_captured_when_present(tmp_path):
    path = tmp_path / "tracker.xlsx"
    _write_tracker(path, [_tracker_row(1, "Pikachu", condition="Near Mint")])
    cards = read_inventory(path)
    assert cards[0].condition == "Near Mint"


def test_blank_rows_are_skipped(tmp_path):
    path = tmp_path / "tracker.xlsx"
    _write_tracker(path, [
        _tracker_row(1, "Pikachu"),
        [None] * 10,
        _tracker_row(2, "Charizard", card_no="4/102"),
    ])
    cards = read_inventory(path)
    assert len(cards) == 2


def test_missing_required_field_raises_with_row_and_column(tmp_path):
    path = tmp_path / "tracker.xlsx"
    _write_tracker(path, [
        _tracker_row(1, "Pikachu"),
        _tracker_row(2, "", card_no="4/102"),
    ])
    with pytest.raises(CardValidationError) as excinfo:
        read_inventory(path)
    assert excinfo.value.row == 6
    assert excinfo.value.column == "Card Name"


def test_notes_without_bracket_tag_raises(tmp_path):
    path = tmp_path / "tracker.xlsx"
    _write_tracker(path, [_tracker_row(1, "Pikachu", notes="no bracket tag here")])
    with pytest.raises(CardValidationError) as excinfo:
        read_inventory(path)
    assert excinfo.value.column == "Notes"


def test_type_mismatch_with_notes_tag_raises(tmp_path):
    path = tmp_path / "tracker.xlsx"
    _write_tracker(path, [_tracker_row(1, "Pikachu", card_type="Lot", notes="[SINGLE] note")])
    with pytest.raises(CardValidationError) as excinfo:
        read_inventory(path)
    assert excinfo.value.column == "Type"


def test_duplicate_card_id_raises(tmp_path):
    path = tmp_path / "tracker.xlsx"
    _write_tracker(path, [
        _tracker_row(1, "Pikachu"),
        _tracker_row(1, "Pikachu Reprint", card_no="2/2"),
    ])
    with pytest.raises(CardValidationError) as excinfo:
        read_inventory(path)
    assert excinfo.value.column == "ID"


def _write_sale_plan(tracker_path, sale_plan_rows):
    workbook = openpyxl.load_workbook(tracker_path)
    sheet = workbook.create_sheet("Sale Plan")
    sheet.append([])
    sheet.append([None, "Sale Plan"])
    sheet.append([None, "Notes text"])
    sheet.append([])
    sheet.append([None, "#", "Listing", "Cards", "List at", "Floor", "Notes"])
    for row in sale_plan_rows:
        sheet.append([None] + list(row))
    workbook.save(tracker_path)


def test_read_sale_plan_parses_rows(tmp_path):
    path = tmp_path / "tracker.xlsx"
    _write_tracker(path, [_tracker_row(1, "Pikachu")])
    _write_sale_plan(path, [(1, "Pikachu 58/102", 1, 45, 35, "note")])
    rows = read_sale_plan(path)
    assert len(rows) == 1
    assert rows[0].number == 1
    assert rows[0].listing_text == "Pikachu 58/102"
    assert str(rows[0].list_price) == "45"


def test_read_sale_plan_missing_price_raises(tmp_path):
    path = tmp_path / "tracker.xlsx"
    _write_tracker(path, [_tracker_row(1, "Pikachu")])
    _write_sale_plan(path, [(1, "Pikachu 58/102", 1, None, 35, "note")])
    with pytest.raises(CardValidationError):
        read_sale_plan(path)
